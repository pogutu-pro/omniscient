"""The trimester calendar, and the timetable importer that reads real sheets.

The importer is exercised against a workbook built here rather than the real
departmental spreadsheet: the real one is a draft that gets revised, and a
test that asserted on its exact contents would fail every time the department
corrected a typo - while still passing if the parser quietly stopped reading
something. Building the sheet keeps the test about the parser.
"""
from __future__ import annotations

import datetime as dt

import pytest
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.academics import AcademicTerm, Programme
from app.repositories.academic_repository import SqlAcademicRepository
from app.schemas.academics import AcademicTermCreate
from app.services.term_calendar import (
    academic_year_for,
    current_trimester,
    parse_year_group,
    term_period,
    trimester_for_month,
    trimester_from_label,
)
from app.services.timetable_import import import_timetable, parse_timetable_workbook

# --- The rule itself ---


@pytest.mark.parametrize(
    ("month", "expected"),
    [(1, "Semester 1"), (4, "Semester 1"), (5, "Semester 2"), (8, "Semester 2"), (9, "Semester 3"), (12, "Semester 3")],
)
def test_dekut_year_has_three_trimesters(month: int, expected: str) -> None:
    """Semester 1 Jan-Apr, Semester 2 May-Aug, Semester 3 Sep-Dec."""
    assert f"Semester {trimester_for_month(month)}" == expected


def test_semester_3_falls_in_the_earlier_calendar_year() -> None:
    """The year is named for the September it opens in, so Semester 3 of
    2026/2027 is September-December *2026*, not 2027."""
    assert term_period("2026/2027", 3) == (dt.date(2026, 9, 1), dt.date(2026, 12, 31))
    assert term_period("2026/2027", 1) == (dt.date(2027, 1, 1), dt.date(2027, 4, 30))
    assert term_period("2026/2027", 2) == (dt.date(2027, 5, 1), dt.date(2027, 8, 31))


def test_the_year_opens_in_september() -> None:
    assert academic_year_for(dt.date(2026, 9, 1)) == "2026/2027"
    assert academic_year_for(dt.date(2026, 8, 31)) == "2025/2026"
    assert academic_year_for(dt.date(2027, 1, 15)) == "2026/2027"
    assert current_trimester(dt.date(2026, 9, 26)) == ("2026/2027", 3)
    assert current_trimester(dt.date(2027, 2, 10)) == ("2026/2027", 1)


def test_a_fourth_semester_does_not_exist() -> None:
    with pytest.raises(ValueError, match="three trimesters"):
        term_period("2026/2027", 4)


@pytest.mark.parametrize(
    ("label", "expected"),
    [
        ("SEPTEMBER - DECEMBER 2026", 3),
        ("JANUARY-APRIL 2027", 1),
        ("MAY - AUGUST 2027", 2),
        ("Semester 2", 2),
        ("Trimester 3", 3),
        ("Reporting", None),
    ],
)
def test_a_term_title_resolves_to_a_trimester(label: str, expected: int | None) -> None:
    assert trimester_from_label(label, "2026/2027") == expected


def test_year_group_labels_split_into_year_and_semester() -> None:
    assert parse_year_group("4.2") == (4, 2)
    assert parse_year_group("1.1") == (1, 1)
    # Not a cohort label, or not a cohort that exists: rejected, not guessed.
    assert parse_year_group("YEAR 4") is None
    assert parse_year_group("4.4") is None


# --- The calendar as stored, read back through the repository ---


async def test_calendar_is_derivable_without_any_stored_rows(db_session: AsyncSession) -> None:
    """An empty table must still answer, because the periods are a rule."""
    repo = SqlAcademicRepository(db_session)
    terms = await repo.list_terms()
    assert {t.trimester for t in terms} == {1, 2, 3}
    assert all(t.provisional for t in terms), "generated terms must not look confirmed"
    assert all(t.reporting_date is None for t in terms), "a reporting date is announced, not derived"


async def test_current_term_follows_the_date(db_session: AsyncSession) -> None:
    repo = SqlAcademicRepository(db_session)
    assert (await repo.get_current_term(dt.date(2026, 9, 26))).trimester == 3
    assert (await repo.get_current_term(dt.date(2027, 3, 2))).trimester == 1
    assert (await repo.get_current_term(dt.date(2027, 6, 2))).trimester == 2


async def test_correcting_a_term_updates_it_in_place(db_session: AsyncSession) -> None:
    """When a date slips - industrial action, say - there must be one row,
    not a second one contradicting the first."""
    repo = SqlAcademicRepository(db_session)
    first = await repo.upsert_term(
        AcademicTermCreate(academic_year="2026/2027", trimester=3, start_date="2026-09-01", end_date="2026-12-31")
    )
    corrected = await repo.upsert_term(
        AcademicTermCreate(
            academic_year="2026/2027",
            trimester=3,
            start_date="2026-10-05",
            end_date="2027-01-16",
            reporting_date="2026-10-05",
            provisional=False,
            notes="Postponed by industrial action",
        )
    )
    assert corrected.id == first.id
    assert len(await repo.list_terms("2026/2027")) == 1
    current = await repo.get_current_term(dt.date(2026, 10, 20))
    assert current.start_date == dt.date(2026, 10, 5)
    assert current.provisional is False, "an admin-entered term is no longer provisional"


async def test_a_reporting_date_outside_the_term_is_rejected(db_session: AsyncSession) -> None:
    from pydantic import ValidationError

    with pytest.raises(ValidationError, match="reporting_date"):
        AcademicTermCreate(
            academic_year="2026/2027",
            trimester=3,
            start_date="2026-09-01",
            end_date="2026-12-31",
            reporting_date="2027-02-01",
        )


# --- The importer ---


def _write_sheet(path, *, term_label: str, grid: list[tuple], catalogue: list[tuple]) -> None:
    """A miniature timetable in the same shape the real sheets use.

    `grid` is a flat list of `(day, cohort, sessions)` rows: the department
    gives every day-group pair its own row and merges the day cell down over
    them, and `sessions` is a list of `(text, slots)` starting at column C. A
    blank row separates the grid from the course catalogue underneath.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DEPARTMENT OF COMPUTER SCIENCE "
    sheet["A1"] = f"TEACHING TIMETABLE FOR {term_label}"
    sheet["A2"] = "DEPARTMENT OF COMPUTER SCIENCE"
    sheet["A3"] = "DAY"
    sheet["B3"] = "GROUP"
    for offset, label in enumerate(["7AM-8AM", "8AM-9AM", "9AM-10AM", "10AM-11AM"]):
        sheet.cell(row=3, column=3 + offset, value=label)

    row = 4
    day_run_start: int | None = None
    previous_day: str | None = None
    for day, cohort, sessions in grid:
        if day != previous_day:
            day_run_start = row
            previous_day = day
        sheet.cell(row=row, column=1, value=day)
        sheet.cell(row=row, column=2, value=cohort)
        column = 3
        for text, span in sessions:
            sheet.cell(row=row, column=column, value=text)
            if span > 1:
                sheet.merge_cells(
                    start_row=row, start_column=column, end_row=row, end_column=column + span - 1
                )
            column += span
        row += 1
        # Close the day cell once the next day starts, matching the sheet.
        if day_run_start is not None and row < 4 + len(grid) and grid[row - 4][0] != day:
            sheet.merge_cells(start_row=day_run_start, start_column=1, end_row=row - 1, end_column=1)
            day_run_start = None
    if day_run_start is not None:
        sheet.merge_cells(start_row=day_run_start, start_column=1, end_row=row - 1, end_column=1)

    row += 1  # the blank separator row the real sheets have
    # The catalogue is laid out in the department's own columns: B is the
    # code, C the name, E/F the weekly lecture and lab hours, G the lecturer
    # and H the class size. A row whose first cell is "YEAR n SEMESTER m"
    # starts a new section and applies to the courses beneath it.
    catalogue_columns = {"code": 2, "name": 3, "lecture": 5, "lab": 6, "lecturer": 7, "size": 8}
    for offset, values in enumerate(catalogue):
        if values and values[0] and values[0].upper().startswith("YEAR "):
            sheet.cell(row=row + offset, column=1, value=values[0])
            continue
        code, name, lecture, lab, lecturer, size = values
        for field, value in (
            ("code", code),
            ("name", name),
            ("lecture", lecture),
            ("lab", lab),
            ("lecturer", lecturer),
            ("size", size),
        ):
            if value is not None:
                sheet.cell(row=row + offset, column=catalogue_columns[field], value=value)
    workbook.save(path)


@pytest.fixture
def sample_workbook(tmp_path):
    path = tmp_path / "timetable.xlsx"
    _write_sheet(
        path,
        term_label="SEPTEMBER - DECEMBER 2026",
        grid=[
            ("MONDAY", "YEAR 2.1", [("CCS 2205 - RCL 1", 2), ("SMA 1105 - RC 18", 1)]),
            ("MONDAY", "YEAR 1.1 (CS)", [("SMA 1105 - RC 18", 1)]),
            ("MONDAY", "YEAR 1.1 (FS)", [("SMA 1105 - RCL 1", 1)]),
            ("TUESDAY", "YEAR 2.1", [("CCS 1102 -ONLINE", 1)]),
            ("TUESDAY", "YEAR 1.1 (CS)", []),
            ("TUESDAY", "YEAR 1.1 (FS)", []),
        ],
        catalogue=[
            ["YEAR 2 SEMESTER 1"],
            ("CCS 2205", "Data Structures & Algorithms", 2, 3, "Dr. Kituku", 140),
            ["YEAR 1 SEMESTER 1"],
            ("SMA 1105", "Discrete Mathematics", 3, None, "Dr. Kihuga", 140),
            ("CCS 1102", "Introduction to Cyber Security", 3, None, "Dr. Naivasha", 100),
        ],
    )
    return path


def test_the_importer_reads_the_calendar_and_the_grid(sample_workbook) -> None:
    parsed = parse_timetable_workbook(str(sample_workbook))
    assert parsed.academic_year == "2026/2027"
    assert parsed.term_label == "SEPTEMBER - DECEMBER 2026"
    # September-December is Semester 3 of the three-trimester year, even
    # though the sheet numbers these cohorts "Semester 1".
    assert parsed.official_trimester == 3
    assert {c.code: c.lecturer for c in parsed.courses} == {
        "CCS 2205": "Dr. Kituku",
        "SMA 1105": "Dr. Kihuga",
        "CCS 1102": "Dr. Naivasha",
    }
    catalogue = {c.code: c for c in parsed.courses}
    assert catalogue["CCS 2205"].lecture_hours == 2
    assert catalogue["CCS 2205"].lab_hours == 3
    assert catalogue["CCS 2205"].class_size == 140
    assert catalogue["SMA 1105"].lab_hours is None


def test_a_merged_cell_becomes_one_session_of_the_right_length(sample_workbook) -> None:
    parsed = parse_timetable_workbook(str(sample_workbook))
    monday = {(s.year_group, s.stream, s.course_code): s for s in parsed.sessions if s.day_of_week == 0}
    # Two slots wide: 07:00-09:00.
    merged = monday[("2.1", None, "CCS 2205")]
    assert (merged.start_time, merged.end_time) == ("07:00", "09:00")
    # One slot, laid out after the merged cell, so it starts at 09:00.
    single = monday[("2.1", None, "SMA 1105")]
    assert (single.start_time, single.end_time) == ("09:00", "10:00")
    # The venue comes from the rest of the cell, and "-ONLINE" is a mode.
    assert merged.venue == "RCL 1"
    online = next(s for s in parsed.sessions if s.course_code == "CCS 1102")
    assert online.session_type == "online"


def test_year_one_splits_into_its_two_classes(sample_workbook) -> None:
    parsed = parse_timetable_workbook(str(sample_workbook))
    monday_first_year = {(s.stream, s.course_code) for s in parsed.sessions if s.day_of_week == 0 and s.year_group == "1.1"}
    assert monday_first_year == {("CS", "SMA 1105"), ("FS", "SMA 1105")}


def test_the_sheet_numbering_is_reported_rather_than_silently_rewritten(sample_workbook) -> None:
    """The cohort labels stay as the department printed them; the mismatch
    with the official calendar is surfaced, not resolved behind the user's
    back."""
    parsed = parse_timetable_workbook(str(sample_workbook))
    assert any("three-trimester calendar" in w for w in parsed.warnings)
    assert {c.semester for c in parsed.courses} == {1}


def test_an_unrecognisable_term_title_is_reported(tmp_path) -> None:
    path = tmp_path / "odd.xlsx"
    _write_sheet(
        path,
        term_label="TERM",
        grid=[("MONDAY", "YEAR 2.1", [("CCS 2205 - RCL 1", 2)])],
        catalogue=[["YEAR 2 SEMESTER 1"], ("CCS 2205", "Data Structures", 2, 3, "Dr. Kituku", 140)],
    )
    parsed = parse_timetable_workbook(str(path), "2026/2027")
    assert parsed.official_trimester is None
    assert any("Could not tell which trimester" in w for w in parsed.warnings)


async def test_importing_twice_does_not_duplicate(db_session: AsyncSession, sample_workbook) -> None:
    db_session.add(Programme(code="BCS", name="BSc Computer Science", school="School of Computing & IT"))
    await db_session.commit()

    first = await import_timetable(db_session, parse_timetable_workbook(str(sample_workbook)), "BCS")
    assert first.courses_created == 3
    assert first.sessions_created == 5
    assert first.official_trimester == 3

    # Re-importing an unchanged draft is a no-op, not a second copy of it.
    second = await import_timetable(db_session, parse_timetable_workbook(str(sample_workbook)), "BCS")
    assert (second.courses_created, second.courses_updated) == (0, 0)
    assert (second.sessions_created, second.sessions_updated) == (0, 0)


async def test_a_corrected_draft_updates_rather_than_duplicates(
    db_session: AsyncSession, sample_workbook, tmp_path
) -> None:
    """A room change and a fixed typo land as updates on the rows that already
    exist - the usual case when a department revises a published draft."""
    db_session.add(Programme(code="BCS", name="BSc Computer Science", school="School of Computing & IT"))
    await db_session.commit()
    await import_timetable(db_session, parse_timetable_workbook(str(sample_workbook)), "BCS")

    corrected = tmp_path / "corrected.xlsx"
    _write_sheet(
        corrected,
        term_label="SEPTEMBER - DECEMBER 2026",
        grid=[
            # Same slots, different room.
            ("MONDAY", "YEAR 2.1", [("CCS 2205 - RCL 2", 2), ("SMA 1105 - RC 18", 1)]),
            ("MONDAY", "YEAR 1.1 (CS)", [("SMA 1105 - RC 18", 1)]),
            ("MONDAY", "YEAR 1.1 (FS)", [("SMA 1105 - RCL 1", 1)]),
            ("TUESDAY", "YEAR 2.1", [("CCS 1102 -ONLINE", 1)]),
            ("TUESDAY", "YEAR 1.1 (CS)", []),
            ("TUESDAY", "YEAR 1.1 (FS)", []),
        ],
        catalogue=[
            ["YEAR 2 SEMESTER 1"],
            ("CCS 2205", "Data Structures and Algorithms", 2, 3, "Dr. Kituku", 140),
            ["YEAR 1 SEMESTER 1"],
            ("SMA 1105", "Discrete Mathematics", 3, None, "Dr. Kihuga", 140),
            ("CCS 1102", "Introduction to Cyber Security", 3, None, "Dr. Naivasha", 100),
        ],
    )
    report = await import_timetable(db_session, parse_timetable_workbook(str(corrected)), "BCS")
    assert (report.courses_created, report.courses_updated) == (0, 1), "the corrected course name"
    assert report.sessions_created == 0
    assert report.sessions_updated == 1, "only the session whose room changed"

    monday = {
        (e.course_code, e.start_time): e
        for e in await SqlAcademicRepository(db_session).get_timetable(_all_sessions())
        if e.day_of_week == 0 and e.year_group == "2.1"
    }
    assert monday[("CCS 2205", "07:00")].venue == "RCL 2"
    courses = await SqlAcademicRepository(db_session).list_courses("BCS", None)
    course = next(c for c in courses if c.code == "CCS 2205")
    assert course.name == "Data Structures and Algorithms"


async def test_a_moved_session_needs_prune_to_clear_its_old_slot(
    db_session: AsyncSession, sample_workbook, tmp_path
) -> None:
    """A session that shifts time is a different slot, so the importer adds
    the new one. Without --prune the abandoned slot is left alone on purpose:
    a draft that omits a session is not the same as a department cancelling
    it, and deleting live timetable rows on a guess is not reversible.
    """
    db_session.add(Programme(code="BCS", name="BSc Computer Science", school="School of Computing & IT"))
    await db_session.commit()
    await import_timetable(db_session, parse_timetable_workbook(str(sample_workbook)), "BCS")

    moved = tmp_path / "moved.xlsx"
    _write_sheet(
        moved,
        term_label="SEPTEMBER - DECEMBER 2026",
        grid=[
            # CCS 2205 shifts from 07:00-09:00 to 09:00-11:00.
            ("MONDAY", "YEAR 2.1", [("SMA 1105 - RC 18", 2), ("CCS 2205 - RCL 1", 2)]),
            ("MONDAY", "YEAR 1.1 (CS)", [("SMA 1105 - RC 18", 1)]),
            ("MONDAY", "YEAR 1.1 (FS)", [("SMA 1105 - RCL 1", 1)]),
            ("TUESDAY", "YEAR 2.1", [("CCS 1102 -ONLINE", 1)]),
            ("TUESDAY", "YEAR 1.1 (CS)", []),
            ("TUESDAY", "YEAR 1.1 (FS)", []),
        ],
        catalogue=[
            ["YEAR 2 SEMESTER 1"],
            ("CCS 2205", "Data Structures & Algorithms", 2, 3, "Dr. Kituku", 140),
            ["YEAR 1 SEMESTER 1"],
            ("SMA 1105", "Discrete Mathematics", 3, None, "Dr. Kihuga", 140),
            ("CCS 1102", "Introduction to Cyber Security", 3, None, "Dr. Naivasha", 100),
        ],
    )
    unpruned = await import_timetable(db_session, parse_timetable_workbook(str(moved)), "BCS")
    assert unpruned.sessions_created == 2
    assert unpruned.sessions_removed == 0, "nothing is deleted without being asked"

    pruned = await import_timetable(
        db_session, parse_timetable_workbook(str(moved)), "BCS", prune=True
    )
    assert pruned.sessions_removed == 2, "the two abandoned year 2 slots"
    year_two = [
        (e.course_code, e.start_time, e.end_time)
        for e in await SqlAcademicRepository(db_session).get_timetable(_all_sessions())
        if e.year_group == "2.1" and e.day_of_week == 0
    ]
    # SMA 1105 takes 07:00-09:00, so CCS 2205 follows it at 09:00-11:00.
    assert year_two == [("SMA 1105", "07:00", "09:00"), ("CCS 2205", "09:00", "11:00")]


async def test_pruning_removes_only_what_the_sheet_dropped(
    db_session: AsyncSession, sample_workbook, tmp_path
) -> None:
    db_session.add(Programme(code="BCS", name="BSc Computer Science", school="School of Computing & IT"))
    await db_session.commit()
    await import_timetable(db_session, parse_timetable_workbook(str(sample_workbook)), "BCS")

    # A revision that still lists year 2, but only one of its two Monday
    # sessions, and no longer mentions year 1 at all.
    revised = tmp_path / "revised.xlsx"
    _write_sheet(
        revised,
        term_label="SEPTEMBER - DECEMBER 2026",
        grid=[("MONDAY", "YEAR 2.1", [("CCS 2205 - RCL 1", 2)])],
        catalogue=[
            ["YEAR 2 SEMESTER 1"],
            ("CCS 2205", "Data Structures & Algorithms", 2, 3, "Dr. Kituku", 140),
        ],
    )
    report = await import_timetable(
        db_session, parse_timetable_workbook(str(revised)), "BCS", prune=True
    )
    # The dropped year 2 sessions go...
    assert report.sessions_removed == 2
    remaining = await SqlAcademicRepository(db_session).get_timetable(_all_sessions())
    assert sorted((e.year_group, e.stream or "", e.course_code) for e in remaining) == [
        # Year 1 is left alone, because a sheet that says nothing about a
        # year group has not said that year group is cancelled...
        ("1.1", "CS", "SMA 1105"),
        ("1.1", "FS", "SMA 1105"),
        # ...and the one year 2 session the revision still lists survives.
        ("2.1", "", "CCS 2205"),
    ]


def _all_sessions():
    from app.schemas.academics import TimetableQuery

    return TimetableQuery(academic_year="2026/2027")


async def test_an_unknown_programme_is_refused_rather_than_invented(
    db_session: AsyncSession, sample_workbook
) -> None:
    from app.services.timetable_import import TimetableImportError

    with pytest.raises(TimetableImportError, match="programme"):
        await import_timetable(db_session, parse_timetable_workbook(str(sample_workbook)), "ZZZ")


async def test_stored_terms_win_over_the_derived_calendar(db_session: AsyncSession) -> None:
    """Once an admin has entered dates, the repository must not answer from
    the rule any more."""
    db_session.add(
        AcademicTerm(
            academic_year="2026/2027",
            trimester=3,
            label="Semester 3",
            start_date=dt.date(2026, 10, 5),
            end_date=dt.date(2027, 1, 16),
            provisional=False,
            notes="Postponed",
        )
    )
    await db_session.commit()
    current = await SqlAcademicRepository(db_session).get_current_term(dt.date(2026, 10, 20))
    assert current.start_date == dt.date(2026, 10, 5)
    assert current.notes == "Postponed"
